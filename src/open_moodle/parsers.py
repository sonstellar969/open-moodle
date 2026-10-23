"""Extract and format text from various course document types."""

from __future__ import annotations

import io
import re
from typing import Optional


def parse_pdf(data: bytes, max_pages: int = 50) -> str:
    """Extract plain text from PDF bytes using PyMuPDF."""
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        return f"[Failed to parse PDF document: {e}]"

    total_pages = len(doc)
    pages_to_read = min(total_pages, max_pages)

    output = []
    output.append(f"# PDF Document ({total_pages} pages total)\n")

    for i in range(pages_to_read):
        page = doc[i]
        text = page.get_text().strip()
        output.append(f"## Page {i + 1}\n")
        if text:
            output.append(text)
        else:
            output.append("[No extractable text on this page - likely an image or scan]")
        output.append("\n")

    if total_pages > max_pages:
        output.append(
            f"\n[Note: Document truncated at {max_pages} pages out of {total_pages}. "
            "Pass a higher page limit to read further.]\n"
        )

    doc.close()
    return "\n".join(output)


def parse_docx(data: bytes) -> str:
    """Extract structured text from DOCX bytes using python-docx."""
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        return f"[Failed to parse Word document: {e}]"

    output = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name.lower() if para.style else ""
        if "heading 1" in style_name:
            output.append(f"# {text}")
        elif "heading 2" in style_name:
            output.append(f"## {text}")
        elif "heading 3" in style_name:
            output.append(f"### {text}")
        elif "list" in style_name:
            output.append(f"- {text}")
        else:
            output.append(text)

    # Process tables if present
    for table_idx, table in enumerate(doc.tables):
        output.append(f"\n### Table {table_idx + 1}\n")
        rows = []
        for row in table.rows:
            cell_texts = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cell_texts))
        if rows:
            output.append(" | ".join(["---"] * len(table.columns)))
            output.extend(rows)
            output.append("\n")

    return "\n\n".join(output)


def parse_pptx(data: bytes, max_slides: int = 60) -> str:
    """Extract slide text and tables from PPTX bytes using python-pptx."""
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        return f"[Failed to parse presentation: {e}]"

    total_slides = len(prs.slides)
    slides_to_read = min(total_slides, max_slides)

    output = []
    output.append(f"# Presentation ({total_slides} slides total)\n")

    for i in range(slides_to_read):
        slide = prs.slides[i]
        slide_lines = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_lines.append(line)
            elif shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_texts = [c.text.strip().replace("\n", " ") for c in row.cells]
                    slide_lines.append(" | ".join(row_texts))

        output.append(f"## Slide {i + 1}")
        if slide_lines:
            output.append("\n".join(slide_lines))
        else:
            output.append("[Empty slide or graphic-only]")
        output.append("\n")

    if total_slides > max_slides:
        output.append(
            f"\n[Note: Presentation truncated at {max_slides} slides out of {total_slides}.]\n"
        )

    return "\n".join(output)


def parse_text(data: bytes) -> str:
    """Decode raw bytes into plain text with multiple encoding fallbacks."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def strip_html(html_str: str) -> str:
    """Clean basic HTML tags into plain text."""
    clean = re.sub(r"<br\s*/?>", "\n", html_str)
    clean = re.sub(r"</p>", "\n\n", clean)
    clean = re.sub(r"<[^>]+>", "", clean)
    # Decode common HTML entities
    clean = clean.replace("&nbsp;", " ").replace("&amp;", "&").replace("&quot;", '"')
    clean = clean.replace("&lt;", "<").replace("&gt;", ">")
    return re.sub(r"\n{3,}", "\n\n", clean).strip()


def extract_content(
    data: bytes,
    filename: str = "",
    mime_type: str = "",
    max_pages: int = 50,
) -> str:
    """Inspect file signature and name, then route to the appropriate parser."""
    import urllib.parse
    import zipfile

    # Extract clean file extension even from URLs with query strings
    raw_name = filename.split("?")[0].split("#")[0]
    clean_name = urllib.parse.unquote(raw_name.split("/")[-1]).lower()
    mime = mime_type.lower()

    # 1. Magic bytes detection
    if data.startswith(b"%PDF-"):
        return parse_pdf(data, max_pages=max_pages)

    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = zf.namelist()
                if any("word/document.xml" in n for n in names):
                    return parse_docx(data)
                if any("ppt/presentation.xml" in n for n in names):
                    return parse_pptx(data, max_slides=max_pages)
        except Exception:
            pass

    # 2. Filename extension detection
    if clean_name.endswith(".pdf") or "pdf" in mime:
        return parse_pdf(data, max_pages=max_pages)

    if clean_name.endswith(".docx") or "wordprocessingml" in mime:
        return parse_docx(data)

    if clean_name.endswith(".pptx") or "presentationml" in mime:
        return parse_pptx(data, max_slides=max_pages)

    if clean_name.endswith((".txt", ".md", ".py", ".c", ".cpp", ".java", ".json", ".csv", ".yaml", ".yml")):
        return parse_text(data)

    if clean_name.endswith((".html", ".htm")) or "html" in mime:
        return strip_html(parse_text(data))

    # 3. Fallback to plain text attempt
    try:
        sample = data[:1024].decode("utf-8")
        if sum(1 for c in sample if c.isprintable() or c.isspace()) / max(len(sample), 1) > 0.85:
            return parse_text(data)
    except Exception:
        pass

    return f"[Binary file '{filename}' of type '{mime_type}' ({len(data)} bytes) cannot be converted to plain text.]"
